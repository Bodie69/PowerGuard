from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL

doc = Document()
doc.add_heading('PowerGuard IoT - Visual System Map', 0)
doc.add_paragraph('This diagram illustrates the physical and logical flow of data across the PowerGuard system.')

table = doc.add_table(rows=7, cols=3)
table.style = 'Table Grid'

def set_cell(r, c, text, bold=False):
    cell = table.cell(r, c)
    cell.text = text
    p = cell.paragraphs[0]
    p.alignment = 1 # Center
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if bold and p.runs:
        p.runs[0].bold = True

# Row 0: Hardware Layer
set_cell(0, 0, 'PZEM-004T Sensors\n(Edge Devices)', True)
set_cell(0, 1, '----(Modbus RTU)---->')
set_cell(0, 2, 'Raspberry Pi Pico W\n(Controller)', True)

# Row 1-2: To Cloud
set_cell(1, 2, '|')
set_cell(2, 2, 'v')

# Row 3: Backend Layer
set_cell(3, 2, 'Supabase Cloud\n(PostgreSQL + Realtime)', True)
set_cell(3, 1, '<----(Polling)----')

# Row 4-5: To Client
set_cell(4, 2, '|')
set_cell(5, 2, 'v')

# Row 6: User Layer
set_cell(6, 2, 'Flutter Mobile App\n(User Interface)', True)
set_cell(6, 1, '<---(WebSockets)---')

doc.add_heading('Key Data Paths:', level=2)
doc.add_paragraph('1. Sensor Data: Voltage/Current/Power values flow into the Pico via UART every 2 seconds.')
doc.add_paragraph('2. Metrics Upload: Pico W pushes JSON data to Supabase using REST API (HTTPS).')
doc.add_paragraph('3. Live Dashboard: Supabase broadcasts database changes to the Flutter app instantly.')
doc.add_paragraph('4. Command Loop: Admin actions (OFF command) flow from App to Supabase, then polled by Pico.')

output_path = 'C:/Users/benf/AppData/Local/Google/AndroidStudio2026.1.1/projects/powerguard_iot.e8985eb4/.artifacts/20260829-204335-a85a3ab0-df18-4aa2-87d9-8b239903b10a/PowerGuard_Visual_Map.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
